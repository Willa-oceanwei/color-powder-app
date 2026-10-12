from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from utils.customer_repository import CustomerInput, create_customer
from utils.database import DatabaseConfig, initialize_database
from utils.recipe_repository import create_recipe
from utils.shipment_catalog import complete_items, shipment_choices
from utils.shipment_codes import printed_code
from utils.shipment_repository import ShipmentError, blank_shipment, get_shipment, printable_shipment, save_shipment, shipment_price_history


@pytest.fixture
def config(tmp_path):
    config = DatabaseConfig(backend="sqlite", path=tmp_path / "display.db")
    initialize_database(config.path)
    create_customer(config, CustomerInput("C1", "Customer"))
    create_recipe(config, {"配方編號": "C0757", "顏色": "Blue", "客戶編號": "C1", "計量單位": "KG"})
    return config


def document():
    doc = dict(blank_shipment(), customer_id="C1", customer_name="Customer")
    doc["items"] = [dict(code="C0757", display_code="0757", name="Blue", unit="KG", quantity="2", price="10")]
    return doc


def test_display_persists_without_changing_price_identity(config):
    saved = save_shipment(config, document(), sync_products=True)
    item = get_shipment(config, saved["id"])["items"][0]
    assert item["code"] == "C0757" and item["display_code"] == "0757"
    html = printable_shipment(saved)
    assert '<td class="code">0757</td>' in html
    assert '<td class="code">C0757</td>' not in html
    history = shipment_price_history(config, "C1", saved["shipment_date"])
    result = complete_items(config, saved, [dict(code="0757", name="", unit="", price="", quantity="1")], [], history=history)
    assert result[0]["code"] == "C0757" and result[0]["price"] == "10"
    assert "0757" not in shipment_choices(config, "C1", all_recipes=True)


def test_invalid_display_and_changed_product(config):
    doc = document()
    doc["items"][0]["display_code"] = "9999"
    with pytest.raises(ShipmentError, match="顯示編號"):
        save_shipment(config, doc)
    assert printed_code(doc["items"][0]) == "C0757"
    create_recipe(config, {"配方編號": "C0863A", "顏色": "Red", "客戶編號": "C1", "計量單位": "KG"})
    old = document()["items"]
    changed = dict(old[0], code="C0863A")
    result = complete_items(config, document(), [changed], old, history=[])
    assert not result[0]["display_code"]


def test_statement_and_excel_use_sales_display(config):
    from utils.receivable_report import build_statements, statement_pages, render_statement_print
    from utils.accounting_export import statement_excel
    saved = save_shipment(config, document())
    statements = build_statements(config, saved["account_date"], saved["account_date"])
    html = render_statement_print(statement_pages(statements))
    assert "<td>0757</td>" in html
    workbook = load_workbook(BytesIO(statement_excel(statements)))
    assert workbook["交易明細"].cell(2, 5).value == "0757"


def test_editor_display_choice_and_search_print_preserve_draft(config):
    saved = save_shipment(config, document())
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0, {root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite", path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.button(key="shipment_edit").click().run()
    choice = next(w for w in app.selectbox if w.label.startswith("第 1 項"))
    assert choice.value == "0757"
    choice.set_value("C0757").run()
    assert printed_code(app.session_state["shipment_draft"]["items"][0]) == "C0757"
    next(w for w in app.selectbox if w.label.startswith("第 1 項")).set_value("0757").run()
    assert app.session_state["shipment_draft"]["items"][0]["code"] == "C0757"
    next(w for w in app.text_area if w.label == "備註").set_value("Keep draft").run()
    next(b for b in app.button if b.label == "查詢").click().run()
    assert not app.exception
    assert app.session_state["shipment_draft"]["notes"] == "Keep draft"
    assert printed_code(app.session_state["shipment_draft"]["items"][0]) == "0757"
    app.button(key="shipment_search_print_generate").click().run()
    assert not app.exception
    assert app.session_state["shipment_search_print_snapshot"][1] == saved["id"]
    assert app.session_state["shipment_draft"]["notes"] == "Keep draft"
