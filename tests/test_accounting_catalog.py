import ast
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from utils.accounting_export import statement_excel, statement_pdf, ranking_excel, ranking_pdf
from utils.customer_repository import CustomerInput, create_customer
from utils.database import DatabaseConfig, initialize_database
from utils.product_repository import blank_product, get_product, list_products, product_references, product_from_recipe, save_product, set_product_active
from utils.recipe_repository import create_recipe
from utils.receivable_report import build_statements, statement_pages
from utils.shipment_catalog import complete_items, shipment_choices
from utils.shipment_repository import ShipmentError, blank_shipment, list_shipments, save_shipment


@pytest.fixture
def config(tmp_path):
    config = DatabaseConfig(backend="sqlite", path=tmp_path / "catalog.db")
    initialize_database(config.path)
    create_customer(config, CustomerInput("C01", "範例"))
    create_customer(config, CustomerInput("C02", "其他客戶"))
    create_recipe(config, {"配方編號": "R1", "顏色": "藍", "客戶編號": "C01", "計量單位": "包"})
    create_recipe(config, {"配方編號": "R2", "顏色": "白", "客戶編號": "C01", "計量單位": "1000g"})
    return config


def shipment(price="125", **updates):
    data = blank_shipment()
    data.update(customer_id="C01", customer_name="範例", shipment_date="2026-10-01", account_date="2026-10-01")
    data["items"] = [dict(code="R1", name="藍", unit="包", quantity="2", price=price, notes="")]
    data.update(updates)
    return data


def test_recipe_defaults_direct_code_and_price_precedence(config):
    refs, _ = product_references(config)
    assert product_from_recipe(refs[1])["specification"] == "1000g"
    choices = shipment_choices(config, "C01")
    assert choices["R1"]["sales_unit"] == "包"
    assert "R1" not in shipment_choices(config, "C02")
    raw = [dict(code="R1", name="", unit="", quantity=2, price="", notes="")]
    item = complete_items(config, shipment(), raw, [])[0]
    assert item["name"] == "藍" and item["unit"] == "包" and item["price"] == "0"
    assert item["notes"] == ""
    code_only = complete_items(config, shipment(), [dict(code="R1", name="", unit="", quantity="", price="")], [])[0]
    assert code_only["name"] == "藍" and code_only["unit"] == "包" and code_only["price"] == "0"
    assert code_only["quantity"] == ""
    other_customer = complete_items(config, shipment(customer_id="C02"), raw, [])[0]
    assert other_customer["name"] == "藍" and other_customer["unit"] == "包"
    manual_notes = dict(raw[0], notes="客戶指定備註")
    assert complete_items(config, shipment(), [manual_notes], [])[0]["notes"] == "客戶指定備註"
    product = product_from_recipe(refs[0])
    product["standard_price"] = "200"
    saved = save_product(config, product)
    assert complete_items(config, shipment(), raw, [])[0]["price"] == "200"
    history = save_shipment(config, shipment())
    assert complete_items(config, shipment(shipment_date="2026-10-02"), raw, [])[0]["price"] == "125"
    override = dict(raw[0], price="0", name="手動名稱", unit="桶")
    assert complete_items(config, shipment(), [override], [])[0]["price"] == "0"
    assert complete_items(config, shipment(), [override], [])[0]["name"] == "手動名稱"
    assert complete_items(config, shipment(), [item], [item]) == [item]
    changed = dict(item, code="R2")
    assert complete_items(config, shipment(), [changed], [item])[0]["name"] == "白"
    assert complete_items(config, shipment(), [changed], [item])[0]["notes"] == ""


def test_optional_bidirectional_sync_and_atomic_rollback(config):
    save_shipment(config, shipment())
    assert not list_products(config)
    save_shipment(config, shipment("160"), sync_products=True)
    product = get_product(config, "R1")
    assert product["standard_price"] == "0" and product["recipe_id"] == "R1"
    product["standard_price"] = "200"
    product = save_product(config, product)
    assert "R1" in shipment_choices(config, "C02")
    save_shipment(config, shipment("180"), sync_products=True)
    assert get_product(config, "R1")["standard_price"] == "200"
    save_shipment(config, shipment("260", customer_id="C02", customer_name="其他客戶"), sync_products=True)
    from utils.shipment_catalog import item_price
    assert item_price(config, shipment(shipment_date="2026-10-02"), "R1", "包", product) == "180"
    assert item_price(config, shipment(customer_id="C02", shipment_date="2026-10-02"), "R1", "包", product) == "260"
    set_product_active(config, "R1", product["version"], active=False, reason="test")
    before = len(list_shipments(config))
    bad = shipment()
    bad["items"] += [dict(code="NEW", name="新貨品", unit="KG", quantity="1", price="1"),
                      dict(code="R1", name="藍", unit="桶", quantity="1", price="2")]
    with pytest.raises(ShipmentError, match="已刪除貨品"):
        save_shipment(config, bad, sync_products=True)
    assert len(list_shipments(config)) == before and len(list_products(config, include_inactive=True)) == 1
    bad = shipment()
    bad["items"][0]["unit"] = "桶"
    with pytest.raises(ShipmentError, match="已刪除貨品"):
        save_shipment(config, bad, sync_products=True)
    assert len(list_shipments(config)) == before


def test_unit_change_refreshes_customer_history_price(config):
    history = shipment("350", tax_mode="免稅", tax_rate="0")
    history["items"][0]["unit"] = "桶"
    save_shipment(config, history)
    document = shipment(shipment_date="2026-10-02")
    old = dict(code="R1", name="手動品名", unit="包", quantity="1", price="0", notes="手動附註")
    changed = dict(old, unit="桶")
    completed = complete_items(config, document, [changed], [old])[0]
    assert completed["price"] == "350"
    assert completed["name"] == "手動品名" and completed["notes"] == "手動附註"
    override = dict(changed, price="400")
    assert complete_items(config, document, [override], [old])[0]["price"] == "400"
    assert complete_items(config, shipment(customer_id="C02"), [changed], [old])[0]["price"] == "0"


def test_customer_sales_name_unit_and_price_do_not_change_recipe(config):
    from utils.recipe_repository import list_recipes
    from utils.shipment_repository import void_shipment
    first = shipment('350')
    first['items'][0].update(name='對外腮紅', unit='g')
    save_shipment(config, first)
    other = shipment('750', customer_id='C02', customer_name='其他客戶')
    other['items'][0].update(name='其他名稱', unit='桶')
    save_shipment(config, other)
    raw = [dict(code='r1', name='', unit='', price='', quantity=1)]
    item = complete_items(config, shipment(), raw, [])[0]
    assert (item['name'], item['unit'], item['price']) == ('對外腮紅', 'g', '350')
    item = complete_items(config, other, raw, [])[0]
    assert (item['name'], item['unit'], item['price']) == ('其他名稱', '桶', '750')
    recipe = next(row for row in list_recipes(config) if row['配方編號'] == 'R1')
    assert recipe['顏色'] == '藍'
    assert recipe['計量單位'] == '包'
    future = shipment('900', shipment_date='2026-10-02')
    future['items'][0].update(name='未來名稱', unit='桶')
    saved = save_shipment(config, future)
    assert complete_items(config, shipment(), raw, [])[0]['unit'] == 'g'
    void_shipment(config, saved['id'], saved['version'], '誤建')
    assert complete_items(config, shipment(shipment_date='2026-10-03'), raw, [])[0]['name'] == '對外腮紅'


def test_color_powder_code_defaults_and_master_precedence(config):
    from utils.color_powder_repository import ColorPowderInput, create_color_powder, set_color_powder_active
    create_color_powder(config, ColorPowderInput('6771', name='色粉黃', package='25KG'))
    raw = [dict(code='6771', name='', unit='', price='', quantity=1)]
    item = complete_items(config, shipment(), raw, [])[0]
    assert item['name'] == '色粉黃' and item['unit'] == 'KG' and item['price'] == '0'
    first = shipment('6')
    first['items'] = [dict(item, name='對外黃', unit='g', price='6')]
    save_shipment(config, first)
    assert complete_items(config, shipment(), raw, [])[0]['unit'] == 'g'
    create_recipe(config, {'配方編號': '6771', '顏色': '配方黃', '計量單位': '包'})
    assert shipment_choices(config, 'C01')['6771']['name'] == '配方黃'
    master = blank_product()
    master.update(product_id='6771', name='貨品黃')
    saved = save_product(config, master)
    assert shipment_choices(config, 'C01')['6771']['name'] == '貨品黃'
    set_product_active(config, '6771', saved['version'], active=False, reason='test')
    assert '6771' not in shipment_choices(config, 'C01')
    create_color_powder(config, ColorPowderInput('DISABLED', name='不使用'))
    set_color_powder_active(config, 'DISABLED', active=False, reason='test')
    assert 'DISABLED' not in shipment_choices(config, 'C01')


def test_add_recipe_with_manually_selected_sales_unit_uses_history(config):
    history = shipment("350", tax_mode="免稅", tax_rate="0")
    history["items"][0]["unit"] = "桶"
    save_shipment(config, history)
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0,{root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.button(key="shipment_new").click().run()
    prefix = f"shipment_{app.session_state['shipment_epoch']}_"
    app.selectbox(key=prefix + "customer_picker").set_value("C01").run()
    from datetime import date
    app.date_input(key=prefix + "shipment_date").set_value(date(2026, 10, 2)).run()
    app.selectbox(key=prefix + "recipe_picker_C01").set_value("R1").run()
    assert app.text_input(key=prefix + "sale_unit").value == "桶"
    app.text_input(key=prefix + "sale_unit").set_value("桶").run()
    app.button(key="shipment_add_recipe").click().run()
    assert not app.exception
    item = app.session_state["shipment_draft"]["items"][0]
    assert item["unit"] == "桶" and str(item["price"]) in ("350", "350.0")
    assert item["notes"] == ""
    app.selectbox(key=prefix + "tax_rate").set_value(0).run()
    assert app.session_state["shipment_draft"]["tax_mode"] == "免稅"
    assert str(app.session_state["shipment_draft"]["items"][0]["price"]) in ("350", "350.0")


@pytest.mark.parametrize("value,expected", [("260.000", "260"), ("260.500", "260.5"), ("0.000", "0"), ("1000000.125", "1000000.125")])
def test_display_price(value, expected):
    from utils.shipment_print import display_price
    assert display_price(value) == expected


def test_grid_code_only_submission_and_save_validation(config):
    import json
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0,{root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.button(key="shipment_new").click().run()
    prefix = f"shipment_{app.session_state['shipment_epoch']}_"
    app.selectbox(key=prefix + "customer_picker").set_value("C01").run()
    columns = json.loads(app.dataframe[0].proto.columns)
    assert not columns["數量"].get("required", False)
    assert not columns["單價"].get("required", False)
    app.session_state[prefix + "items"] = {"edited_rows": {}, "deleted_rows": [], "added_rows": [{"貨品編號": "R1"}]}
    app.run()
    assert not app.exception
    item = app.session_state["shipment_draft"]["items"][0]
    assert item["name"] == "藍" and item["unit"] == "包" and item["quantity"] == ""
    app.button(key="shipment_save").click().run()
    assert not app.exception and app.error and not list_shipments(config)


def test_tax_choices_sync_with_invoice_mode(config):
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0,{root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.button(key="shipment_new").click().run()
    epoch = app.session_state["shipment_epoch"]
    key = f"shipment_{epoch}_"
    app.selectbox(key=key + "tax_rate").set_value(0).run()
    assert not app.exception
    assert app.session_state["shipment_draft"]["tax_mode"] == "免稅"
    assert app.session_state["shipment_draft"]["tax_rate"] == "0"
    app.selectbox(key=key + "tax_rate").set_value(5).run()
    assert not app.exception and app.session_state["shipment_draft"]["tax_mode"] == "外加"
    app.selectbox(key=key + "tax_mode").set_value("內含").run()
    assert not app.exception and app.session_state["shipment_draft"]["tax_rate"] == "5"
    app.selectbox(key=key + "tax_mode").set_value("免稅").run()
    assert not app.exception and app.session_state["shipment_draft"]["tax_rate"] == "0"


def test_invoice_dropdowns_tax_autofill_and_manual_override(config):
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0,{root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.button(key="shipment_new").click().run()
    prefix = f"shipment_{app.session_state['shipment_epoch']}_"
    app.selectbox(key=prefix + "customer_picker").set_value("C01").run()
    app.session_state[prefix + "items"] = {"edited_rows": {}, "deleted_rows": [], "added_rows": [
        {"貨品編號": "R1", "品名": "藍", "單位": "包", "數量": 1.0, "單價": 1000.0}]}
    app.run()
    assert not app.exception
    assert app.text_input(key=prefix + "invoice_amount").value == "50"
    assert app.selectbox(key=prefix + "invoice_method").value == "隨單開立"
    assert app.selectbox(key=prefix + "invoice_type").value == "三聯式"
    app.selectbox(key=prefix + "tax_rate").set_value(0).run()
    assert app.text_input(key=prefix + "invoice_amount").value == "0"
    app.selectbox(key=prefix + "tax_rate").set_value(5).run()
    assert app.text_input(key=prefix + "invoice_amount").value == "50"
    app.text_input(key=prefix + "invoice_amount").set_value("75").run()
    app.selectbox(key=prefix + "invoice_method").set_value("月結開立").run()
    app.selectbox(key=prefix + "invoice_type").set_value("收銀機").run()
    epoch = app.session_state["shipment_grid_epoch"]
    app.session_state[prefix + "items" + (f"_{epoch}" if epoch else "")] = {
        "edited_rows": {0: {"單價": 2000.0}} if epoch else {}, "deleted_rows": [],
        "added_rows": [] if epoch else [{"貨品編號": "R1", "品名": "藍", "單位": "包", "數量": 1.0, "單價": 2000.0}]}
    app.run()
    assert not app.exception and app.text_input(key=prefix + "invoice_amount").value == "75"
    app.button(key="shipment_save").click().run()
    app.run()
    assert not app.exception
    from utils.shipment_repository import get_shipment
    saved = get_shipment(config, app.session_state["shipment_selected"])
    assert saved["tax_amount"] == "100" and saved["total_amount"] == "2100"
    assert saved["invoice"] == dict(date="", number="", method="月結開立", type="收銀機", amount="75")
    app.button(key="shipment_edit").click().run()
    assert not app.exception
    assert "shipment_draft" in app.session_state
    prefix = f"shipment_{app.session_state['shipment_epoch']}_"
    assert app.text_input(key=prefix + "invoice_amount").value == "75"


def test_report_exports_keep_totals_and_text_safe(config):
    doc = shipment()
    doc["items"][0]["name"] = "=HYPERLINK(\"danger\")"
    save_shipment(config, doc)
    statements = build_statements(config, "2026-10-01", "2026-10-31")
    workbook = load_workbook(BytesIO(statement_excel(statements)))
    assert workbook["交易明細"]["F2"].value.startswith("=HYPERLINK")
    assert workbook["交易明細"]["F2"].data_type == "s"
    assert workbook["客戶合計"]["G2"].value == 263
    assert workbook["客戶合計"].max_column == 7
    assert load_workbook(BytesIO(statement_excel(statements, show_receipts=True)))["客戶合計"].max_column == 9
    pdf = statement_pdf(statement_pages(statements))
    assert pdf.startswith(b"%PDF") and b"/UniCNS-UCS2-H" in pdf
    assert b"595.2756 419.5276" in pdf
    assert ranking_pdf(statements).startswith(b"%PDF")
    assert load_workbook(BytesIO(ranking_excel(statements))).active["E2"].value == 250


def test_shipment_editing_reuses_reads_and_direct_grid_autofill(config, monkeypatch):
    import utils.shipment_ui as ui
    calls = []
    for name in ("list_customers", "list_shipments", "shipment_choices", "shipment_price_history"):
        original = getattr(ui, name)
        def wrapped(*args, _name=name, _original=original, **kwargs):
            calls.append(_name)
            return _original(*args, **kwargs)
        monkeypatch.setattr(ui, name, wrapped)
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0,{root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.button(key="shipment_new").click().run()
    prefix = f"shipment_{app.session_state['shipment_epoch']}_"
    app.selectbox(key=prefix + "customer_picker").set_value("C01").run()
    warm = len(calls)
    app.text_input(key=prefix + "address").set_value("送貨地址").run()
    app.text_area(key=prefix + "notes").set_value("備註").run()
    app.selectbox(key=prefix + "tax_rate").set_value(0).run()
    app.session_state[prefix + "items"] = {"edited_rows": {}, "deleted_rows": [], "added_rows": [
        {"貨品編號": "R1", "數量": 2.0}, {"貨品編號": "R2", "數量": 1.0}]}
    app.run()
    assert not app.exception
    assert app.session_state["shipment_draft"]["items"][0]["name"] == "藍"
    assert app.session_state["shipment_draft"]["items"][0]["unit"] == "包"
    import json
    columns = json.loads(app.dataframe[0].proto.columns)
    assert all(not columns[key].get('disabled', False) for key in ('貨品編號', '品名', '單位', '採購單號', '附註說明'))
    assert len(calls) == warm
    app.button(key="shipment_reload_data").click().run()
    assert len(calls) > warm and len(app.session_state["shipment_draft"]["items"]) == 2


def test_product_recipe_ui_and_hidden_old_fields(config):
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0,{root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.product_ui import render_product_management
render_product_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
''').run(timeout=30)
    app.selectbox(key="product_recipe_search").set_value("R1").run()
    app.button(key="product_recipe_open").click().run()
    assert not app.exception
    assert next(w for w in app.text_input if w.label == "規格").value == "包"
    assert next(w for w in app.text_input if w.label == "貨品編號").value == "R1"
    assert not any(w.label.startswith("外包裝") for w in app.text_input)
    assert not any(w.label in ("售價 C", "售價 D", "售價 E") for w in app.number_input)
    next(w for w in app.number_input if w.label == "標準售價").set_value(260.0).run()
    app.button(key="product_save").click().run()
    assert not app.exception and get_product(config, "R1")["standard_price"] == "260.0"


def test_navigation_retains_shipment_grid_and_inputs(config, monkeypatch):
    from streamlit.elements.lib import policies
    check = policies.check_session_state_rules

    def check_every_widget(*args, **kwargs):
        # Streamlit normally shows this warning only once per process.
        policies._shown_default_value_warning = False
        return check(*args, **kwargs)

    monkeypatch.setattr(policies, "check_session_state_rules", check_every_widget)
    root = Path(__file__).resolve().parents[1]
    tree = ast.parse((root / "app.py").read_text(encoding="utf-8"))
    sidebar = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "render_sidebar")
    script = f'''import sys
sys.path.insert(0,{str(root)!r})
from pathlib import Path
import streamlit as st
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
def request_logout(): pass
''' + ast.unparse(sidebar) + f'''
render_sidebar()
if st.session_state.menu == "出貨單":
    render_shipment_management(DatabaseConfig(backend="sqlite",path=Path({str(config.path)!r})))
'''
    app = AppTest.from_string(script).run(timeout=30)
    app.button(key="出貨單").click().run()
    app.button(key="shipment_new").click().run()
    next(w for w in app.selectbox if w.label == "客戶").set_value("C01").run()
    date_widget_id = next(w for w in app.date_input if w.label == "出貨日期").proto.id
    address_widget_id = next(w for w in app.text_input if w.label == "送貨地址").proto.id
    next(w for w in app.text_input if w.label == "送貨地址").set_value("草稿地址").run()
    from datetime import date
    next(w for w in app.date_input if w.label == "出貨日期").set_value(date(2026, 10, 3)).run()
    assert next(w for w in app.date_input if w.label == "出貨日期").proto.id == date_widget_id
    assert next(w for w in app.text_input if w.label == "送貨地址").proto.id == address_widget_id
    assert not any("Session State API" in w.value for w in app.warning)
    epoch = app.session_state["shipment_epoch"]
    app.session_state[f"shipment_{epoch}_items"] = {"edited_rows": {}, "deleted_rows": [], "added_rows": [
        {"貨品編號": "R1", "數量": 1.0}]}
    app.run()
    assert not app.exception
    assert app.session_state["shipment_draft"]["items"][0]["name"] == "藍"
    assert not any("Session State API" in w.value for w in app.warning)
    app.button(key="客戶名單").click().run()
    assert not app.exception and app.session_state["menu"] == "客戶名單"
    app.run()
    app.button(key="出貨單").click().run()
    assert not app.exception
    assert next(w for w in app.text_input if w.label == "送貨地址").value == "草稿地址"
    assert next(w for w in app.date_input if w.label == "出貨日期").value == date(2026, 10, 3)
    assert not any("Session State API" in w.value for w in app.warning)
    assert app.session_state["shipment_draft"]["items"][0]["unit"] == "包"
    next(w for w in app.date_input if w.label == "出貨日期").set_value(None).run()
    assert not app.exception
    assert any("請輸入出貨日期" in w.value for w in app.error)
    assert not list_shipments(config)
