from pathlib import Path

from utils.database import DatabaseConfig, initialize_database
from utils.customer_repository import CustomerInput, create_customer
from utils.shipment_catalog import matching_products, complete_items
from utils.shipment_repository import blank_shipment, save_shipment, shipment_price_history, printable_shipment
from utils.product_repository import blank_product


def catalog():
    return {code: dict(blank_product(), product_id=code, name=name, sales_unit="KG")
            for code, name in (("C0757", "Blue"), ("C0863A", "Red"), ("52743A", "Yellow A"), ("52743B", "Yellow B"))}


def test_code_aliases_keep_zeroes_and_exact_codes_win():
    choices = catalog()
    assert matching_products(choices, "0757") == ["C0757"]
    assert matching_products(choices, "0863a") == ["C0863A"]
    assert matching_products(choices, "52743") == ["52743A", "52743B"]
    choices["0757"] = choices["C0757"]
    assert matching_products(choices, "0757") == ["0757"]


def test_name_lookup_and_manual_name_retention():
    choices = catalog()
    history = [dict(code="C0863A", name="Customer red", unit="g", price="12", tax_mode="外加", tax_exempt=False)]
    doc = dict(blank_shipment(), customer_id="C1")
    row = dict(code="", name="Customer red", quantity="2", unit="", price="")
    result = complete_items(None, doc, [row], [], choices=choices, history=history)
    assert result[0]["code"] == "C0863A"
    assert (result[0]["name"], result[0]["unit"], result[0]["price"]) == ("Customer red", "g", "12")
    edited = dict(result[0], name="Manual customer name")
    assert complete_items(None, doc, [edited], result, choices=choices, history=history)[0]["name"] == "Manual customer name"
    assert choices["C0863A"]["name"] == "Red"
    assert matching_products(choices, "Yellow", by_name=True) == ["52743A", "52743B"]


def test_latest_customer_name_and_unit_use_shipment_date(tmp_path):
    config = DatabaseConfig(backend="sqlite", path=tmp_path / "history.db")
    initialize_database(config.path)
    create_customer(config, CustomerInput("C1", "Customer"))
    for day, name, unit, price in ((5, "Newest name", "g", "12"), (1, "Old name", "KG", "10"), (3, "Middle name", "桶", "11")):
        doc = dict(blank_shipment(), customer_id="C1", customer_name="Customer", shipment_date=f"2026-10-{day:02d}")
        doc["items"] = [dict(code="C0863A", name=name, quantity="2", unit=unit, price=price)]
        save_shipment(config, doc)
    history = shipment_price_history(config, "C1", "2026-10-07")
    result = complete_items(None, dict(blank_shipment(), customer_id="C1"),
                            [dict(code="0863A", name="", quantity="1", unit="", price="")], [], choices=catalog(), history=history)
    assert (result[0]["name"], result[0]["unit"], result[0]["price"]) == ("Newest name", "g", "12")


def test_date_number_can_be_overridden_and_feed_margins(tmp_path):
    config = DatabaseConfig(backend="sqlite", path=tmp_path / "numbers.db")
    initialize_database(config.path)
    create_customer(config, CustomerInput("C1", "Customer"))
    doc = dict(blank_shipment(), customer_id="C1", customer_name="Customer", shipment_number="CUSTOM-001", number_mode="date")
    doc["items"] = [dict(code="C0757", name="Blue", quantity="2", unit="KG", price="10")]
    saved = save_shipment(config, doc)
    assert saved["shipment_number"] == "CUSTOM-001"
    html = printable_shipment(saved, feed_rotation=90, hole_margin=12.5)
    assert "size:A5 portrait" in html and "rotate(90deg)" in html
    assert "padding-top:12.5mm" in html and "bottom:12.5mm" in html
    html = printable_shipment(saved, hole_margin=12.5)
    assert "padding-left:12.5mm" in html and "padding-right:12.5mm" in html


def test_search_during_entry_preserves_customer_and_notes(tmp_path):
    from streamlit.testing.v1 import AppTest
    config = DatabaseConfig(backend="sqlite", path=tmp_path / "search.db")
    initialize_database(config.path)
    create_customer(config, CustomerInput("C1", "One"))
    create_customer(config, CustomerInput("C2", "Two"))
    root = str(Path(__file__).resolve().parents[1])
    app = AppTest.from_string(f'''import sys
sys.path.insert(0, {root!r})
from pathlib import Path
from utils.database import DatabaseConfig
from utils.shipment_ui import render_shipment_management
render_shipment_management(DatabaseConfig(backend="sqlite", path=Path({str(config.path)!r})), start_new_on_entry=True)
''').run(timeout=30)
    prefix = f"shipment_{app.session_state['shipment_epoch']}_"
    assert app.selectbox(key=prefix + "number_mode").value == "自行輸入"
    app.selectbox(key=prefix + "customer_picker").set_value("C2").run()
    app.text_area(key=prefix + "notes").set_value("Keep draft").run()
    app.session_state[prefix + "items"] = {"edited_rows": {}, "deleted_rows": [], "added_rows": [
        {"貨品編號": "P1", "品名": "Manual name", "數量": 2.0, "單位": "g", "單價": 7.0}]}
    app.run()
    items = app.session_state["shipment_draft"]["items"].copy()
    next(b for b in app.button if b.label == "查詢").click().run()
    assert not app.exception
    assert app.session_state["shipment_draft"]["customer_id"] == "C2"
    assert app.session_state["shipment_draft"]["notes"] == "Keep draft"
    assert app.session_state["shipment_draft"]["items"] == items


def test_report_title_totals_and_ranking_export():
    from io import BytesIO
    from openpyxl import load_workbook
    from utils.receivable_report import statement_pages, render_statement_print
    from utils.accounting_export import ranking_excel
    statement = dict(customer_id="C1", customer_name="Customer", documents=[], start="2026-10-01", end="2026-10-12",
                     net="440", tax="22", total="462", received="0", balance="462", quantities={"g": "40"})
    html = render_statement_print(statement_pages([statement]), title="客戶交易明細")
    assert "<title>客戶交易明細</title>" in html and "<h2>客戶交易明細</h2>" in html
    assert "grid-column:6/10" in html and "<span>本期總計：</span><span>462</span>" in html
    book = load_workbook(BytesIO(ranking_excel([statement])))
    assert book.active.cell(3, 3).value == "未稅金額總計"
    assert book.active.cell(3, 5).value == 440
