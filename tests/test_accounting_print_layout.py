from copy import deepcopy
from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import load_workbook

from utils.accounting_export import statement_excel, statement_pdf
from utils.receivable_report import statement_pages, render_statement_print
from utils.shipment_print import render_shipment_print
from utils.print_stock import stock_geometry


def statement():
    item = dict(code="C0757", name="藍", quantity="2", unit="桶", price="100", amount="200",
                order_number="PO-20261001", notes="Not a purchase order")
    return dict(customer_id="C1", customer_name="測試客戶", contact="王先生", phone="06-1234567", fax="06-7654321",
                tax_id="12345678", start="2026-10-01", end="2026-10-31", quantities={"桶": "2"},
                net="200", tax="10", total="210", received="0", balance="210",
                documents=[dict(date="2026-10-01", number="S1", items=[item])])


def test_a4_driver_shared_feed_preserves_data_and_blanks_prices():
    s = statement()
    doc = dict(customer_id=s["customer_id"], customer_name=s["customer_name"], shipment_number="S1",
               shipment_date="2026-10-01", items=s["documents"][0]["items"], net_amount="200", tax_amount="10", total_amount="210")
    before = deepcopy(doc)
    html = render_shipment_print(doc, paper_stock="a4_driver", feed_rotation=90, hole_margin=12.5, show_prices=False)
    assert 'size:A4 portrait' in html and 'width:210mm;height:297mm;' in html
    assert 'rotate(0deg)' in html and 'width:240mm;height:140mm' in html
    assert '<th class="price">單價</th>' in html and '<th class="amount">金額</th>' in html
    assert '<td class="price"></td>' in html and '<td class="amount"></td>' in html
    assert '<span>合計' not in html and doc == before
    assert 'font-size:13pt' in html and 'font-size:20pt' in html
    x = stock_geometry(0, 12.5, offset=-25, preserve_holes=False)[2]
    assert f'translate({x:.6f}mm' in html
    assert stock_geometry(0, 12.5)[2] - x == 25
    assert '.stock-page>.sheet{width:222mm}' in html
    assert '.company-address {font-size:9pt;font-family:"DFKai-SB"' in html
    assert '.company-contact {font-size:11pt;}' in html
    assert '.order_number {padding-left:4mm;}' in html
    assert 'h1 {margin:0;font-size:21pt;' in html
    assert 'margin:2mm 0 .7mm;' in html
    assert 'th {padding-top:.65mm;padding-bottom:.65mm;}' in html
    assert '.metadata {display:grid;grid-template-columns:1fr 56mm' in html
    assert '.delivery-row {display:grid;grid-template-columns:1fr 56mm' in html
    # Enlarged content stays inside the logical A4 width used by the driver.
    scale = stock_geometry(0, 12.5)[4]
    assert x + (222 - 10) * scale < 210


def test_a4_statement_headers_alignment_company_and_purchase_order():
    s = statement()
    before = deepcopy(s)
    pages = statement_pages([s], paper_size="A4")
    html = render_statement_print(pages, paper_stock="a4_driver", title="客戶交易明細")
    assert 'width:240mm;height:280mm' in html and 'size:A4 portrait' in html
    assert '<th>採購單號</th>' in html and 'PO-20261001' in html and 'Not a purchase order' not in html
    assert 'font-size:13pt' in html and 'table{font-size:12pt}' in html
    assert 'th:nth-child(6),th:nth-child(8),th:nth-child(9){text-align:right}' in html
    assert '<footer>' not in html and 'class="company-info"' in html
    assert '.company-info{margin-top:5.8mm' in html
    assert '.statement-left{grid-column:1/6;grid-row:1' in html
    assert '.statement-totals{grid-column:6/10;grid-row:1' in html
    pdf = statement_pdf(pages, paper_stock="a4_driver", title="客戶交易明細")
    assert b'595.2756 841.8898' in pdf
    assert s == before


def test_statement_excel_print_sheet_keeps_company_and_safe_text():
    s = statement()
    s["documents"][0]["items"][0]["name"] = '=HYPERLINK("danger")'
    book = load_workbook(BytesIO(statement_excel([s], title="客戶交易明細")))
    sheet = book["列印版"]
    assert book.active.title == "列印版"
    assert sheet['A1'].value == "佳咊實業有限公司" and sheet['A2'].value == "客戶交易明細"
    assert sheet['A4'].value == "客戶名稱：C1　測試客戶"
    assert sheet['J8'].value == "採購單號" and sheet['J10'].value == "PO-20261001"
    assert sheet['E10'].value.startswith('=HYPERLINK') and sheet['E10'].data_type == 's'
    assert sheet['A14'].value == "台南市永康區中正南路309號"
    assert sheet['C12'].value == 2 and sheet['I12'].value == 200
    assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A4
    assert sheet.page_setup.fitToWidth == 1
    assert book['交易明細']['L2'].value == 'PO-20261001'


def test_twelve_digit_purchase_order_and_mixed_units_share_summary_height():
    s = statement()
    s['documents'][0]['items'][0]['order_number'] = '123456789012'
    s['quantities']['KG'] = '3'
    pages = statement_pages([s], paper_size='A4')
    assert pages[0]['rows'][1][-1] == '123456789012'
    assert len(pages[0]['rows']) == 2
    html = render_statement_print(pages)
    assert 'td:nth-child(10){font-size:9.5pt;white-space:nowrap' in html
    assert '<div class="statement-left"><div class="quantity-summary">' in html
    book = load_workbook(BytesIO(statement_excel([s])))
    sheet = book['列印版']
    assert sheet['J10'].value == '123456789012'
    assert sheet['J10'].font.sz == 11 and not sheet['J10'].alignment.wrap_text
    assert sheet['C11'].value == 2 and sheet['C12'].value == 3
    assert sheet['I11'].value == 200 and sheet['I13'].value == 210
    assert sheet['A14'].value == '台南市永康區中正南路309號'


def test_excel_separates_each_shipment_even_on_the_same_date():
    s = statement()
    second = deepcopy(s['documents'][0])
    second['number'] = 'S2'
    s['documents'].append(second)
    sheet = load_workbook(BytesIO(statement_excel([s])))['列印版']
    assert sheet['C11'].value == 'S2'
    for column in range(1, 11):
        assert sheet.cell(11, column).border.top.style == 'dotted'
        assert sheet.cell(10, column).border.top.style is None


def test_long_purchase_orders_stay_complete_on_one_row_and_fit_exports(monkeypatch):
    import utils.accounting_export as export
    long_order = 'P-2610-447-1-01-EXTRA-12345678901234567890'
    s = statement()
    s['documents'][0]['items'][0]['order_number'] = long_order
    pages = statement_pages([s], paper_size='A4')
    assert len(pages[0]['rows']) == 2 and pages[0]['rows'][1][-1] == long_order
    html = render_statement_print(pages)
    assert f'<span data-purchase-order>{long_order}</span>' in html
    assert 'beforeprint' in html and 'span.scrollWidth' in html
    doc = dict(customer_id=s['customer_id'], customer_name=s['customer_name'], shipment_number='S1',
               shipment_date='2026-10-01', items=s['documents'][0]['items'], net_amount='200', tax_amount='10', total_amount='210')
    for orientation in ('landscape', 'portrait'):
        shipment = render_shipment_print(doc, orientation=orientation)
        assert shipment.count('<tbody><tr>') == 1 and shipment.count('</tr></tbody>') == 1
        assert f'<span data-purchase-order>{long_order}</span>' in shipment
    sheet = load_workbook(BytesIO(statement_excel([s])))['列印版']
    assert sheet['J10'].value == long_order and sheet['J10'].alignment.shrinkToFit
    assert not sheet['J10'].alignment.wrap_text
    pdf = MagicMock()
    monkeypatch.setattr(export, '_canvas', lambda *args, **kwargs: (BytesIO(), pdf))
    # Initialize the CID font used to measure actual PDF text widths.
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont('MSung-Light'))
    statement_pdf(pages)
    assert any(call.args[0] == 'MSung-Light' and call.args[1] < 9.5 for call in pdf.setFont.call_args_list)
    assert any(call.args[2] == long_order for call in pdf.drawString.call_args_list)


def test_receivable_widening_and_smaller_quantity_totals_do_not_change_other_reports(monkeypatch):
    import utils.accounting_export as export
    pages = statement_pages([statement()], paper_size='A4')
    receivable = render_statement_print(pages, paper_stock='a4_driver')
    other = render_statement_print(pages, paper_stock='a4_driver', title='客戶交易明細')
    x = stock_geometry(0, 12.5, height=280, offset=-25, preserve_holes=False)[2]
    assert f'translate({x:.6f}mm' in receivable
    assert '.stock-page>.sheet.a4{width:222mm}' in receivable
    assert '.quantity-summary{font-size:12pt}' in receivable
    assert '.stock-page>.sheet.a4{width:222mm}' not in other
    assert '.quantity-summary{font-size:12pt}' not in other
    for title, size in [('應收帳款明細表', 12), ('客戶交易明細', 13)]:
        sheet = load_workbook(BytesIO(statement_excel([statement()], title=title)))['列印版']
        assert all(sheet[cell].font.sz == size for cell in ['A11', 'C11', 'D11'])
        assert sheet['I11'].font.sz == 13

    canvases = []
    def canvas(title, *, a5=False):
        pdf = MagicMock()
        canvases.append(pdf)
        return BytesIO(), pdf
    monkeypatch.setattr(export, '_canvas', canvas)
    statement_pdf(pages, paper_stock='a4_driver')
    statement_pdf(pages, paper_stock='a4_driver', title='客戶交易明細')
    report, comparison = canvases
    assert report.translate.call_args.args[0] == comparison.translate.call_args.args[0]
    # Monetary totals keep the same right edge as the item amount column.
    amount_x = [call.args[0] for call in report.drawRightString.call_args_list if call.args[2] == '200']
    assert len(amount_x) == 2 and amount_x[0] == amount_x[1]
    assert report.line.call_args_list[0].args[2] > comparison.line.call_args_list[0].args[2]
    statement_pdf(pages)
    centered = canvases[-1]
    assert not centered.translate.called and not centered.scale.called
    assert centered.drawCentredString.call_args_list[0].args[0] == 297.5


def test_shipment_pdf_mode_centers_without_changing_print_calibration_or_values():
    s = statement()
    doc = dict(customer_id=s['customer_id'], customer_name=s['customer_name'], shipment_number='S1',
               shipment_date='2026-10-01', items=s['documents'][0]['items'], net_amount='200', tax_amount='10', total_amount='210')
    before = deepcopy(doc)
    physical = render_shipment_print(doc, paper_stock='a4_driver', hole_margin=12.5)
    pdf = render_shipment_print(doc, paper_stock='a4_driver', hole_margin=12.5, output_mode='pdf')
    assert 'transform:translate(5mm,0mm) rotate(0deg) scale(0.900900901)' in pdf
    assert 'transform:translate(5mm,0mm)' not in physical
    assert '<td class="amount">200</td>' in pdf and doc == before
    hidden = render_shipment_print(doc, paper_stock='a4_driver', output_mode='pdf', show_prices=False)
    assert '<td class="amount"></td>' in hidden and '<td class="price"></td>' in hidden


def test_statement_pdf_download_does_not_apply_printer_alignment():
    from pathlib import Path
    source = (Path(__file__).resolve().parents[1] / 'utils/receivable_ui.py').read_text(encoding='utf-8')
    assert 'statement_pdf(pages, show_receipts=show_receipts, title=title, company_name=company_name)' in source
    assert 'statement_pdf(pages, show_receipts=show_receipts, title=title, **print_settings)' not in source


def test_selected_statement_company_matches_html_excel_and_embedded_pdf_font():
    from reportlab.pdfbase import pdfmetrics
    s = statement()
    before = deepcopy(s)
    pages = statement_pages([s], paper_size='A4')
    assert '<h1>佳咊實業有限公司</h1>' in render_statement_print(pages)
    for company in ['佳咊實業有限公司', '加舜企業有限公司']:
        assert f'<h1>{company}</h1>' in render_statement_print(pages, company_name=company)
        sheet = load_workbook(BytesIO(statement_excel([s], company_name=company)))['列印版']
        assert sheet['A1'].value == company
        pdf = statement_pdf(pages, company_name=company)
        assert pdf.startswith(b'%PDF')
        from utils.accounting_export import _company_pdf_font
        assert all(ord(char) in pdfmetrics.getFont(_company_pdf_font()).face.charToGlyph for char in company)
    assert s == before


def test_pdf_font_does_not_reuse_an_old_registered_subset():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    from utils.accounting_export import _company_pdf_font
    old = TTFont('CompanySansTC', str(Path(__file__).resolve().parents[1] / 'assets/fonts/CompanySansTC.ttf'))
    pdfmetrics.registerFont(old)
    pdf = statement_pdf(statement_pages([statement()], paper_size='A4'), company_name='加舜企業有限公司')
    current = pdfmetrics.getFont(_company_pdf_font())
    assert current is not pdfmetrics.getFont('CompanySansTC')
    assert all(ord(char) in current.face.charToGlyph for char in '加舜企業有限公司')
    assert pdf.startswith(b'%PDF')

