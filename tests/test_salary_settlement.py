from copy import deepcopy
from io import BytesIO
import sqlite3

import pytest
from openpyxl import load_workbook

from utils.database import DatabaseConfig, connect_from_config, initialize_database_with_health
from utils.salary_calculator import calculate_salary
from utils.salary_excel import generate_salary_workbook, salary_report_signature
from utils.salary_repository import (
    get_salary_monthly_extras, get_settled_month_salaries, save_employee,
    save_salary, settle_salary_month,
)


def _salary(employee_id="E1"):
    adjustments = [
        {"type": "addition", "item_name": "Bonus", "amount": 1800, "note": "Current bonus"},
        {"type": "deduction", "item_name": "Advance", "amount": 500, "note": "Current advance"},
    ]
    salary = {
        "employee_id": employee_id, "employee_name_snapshot": employee_id,
        "base_salary_snapshot": 32900, "attendance_bonus_snapshot": 2000,
        "cooling_allowance_snapshot": 500, "allowance_snapshot": 1000,
        "position_allowance_snapshot": 300, "insurance_snapshot": 1196,
        "standard_hours_snapshot": 8, "leave_days": 1, "leave_hours": 2,
        "annual_leave_days": 1, "annual_leave_hours": 0,
        "annual_leave_balance_before": 10, "annual_leave_balance_after": 9,
        "late_deduction": 200, "system_note": "Edited leave note",
        "manual_note": "Edited manual note", "adjustments": adjustments,
        "annual_leave_records": [{"date": "2026-09-03", "days": 1, "hours": 0, "note": "Leave"}],
    }
    salary.update(calculate_salary(salary, adjustments[:1], adjustments[1:]))
    return salary


def _extras():
    return {"employee_values": {"E1": 125}, "previous_value": 14028,
            "monthly_addition": 125, "monthly_total": 14153}


def test_settlement_download_matches_every_preview_cell(tmp_path):
    config = DatabaseConfig("sqlite", tmp_path / "payroll.db")
    initialize_database_with_health(config)
    salary = _salary()
    save_employee(config, {"employee_id": "E1", "name": "E1"})
    save_salary(config, {**salary, "year": 2026, "month": 9, "base_salary_snapshot": 100})
    preview = load_workbook(BytesIO(generate_salary_workbook(2026, 9, [salary], _extras()))).active

    settle_salary_month(config, 2026, 9, [salary], _extras())

    persisted = get_settled_month_salaries(config, 2026, 9)
    extras = get_salary_monthly_extras(config, 2026, 9)
    downloaded = load_workbook(BytesIO(generate_salary_workbook(2026, 9, persisted, extras))).active
    assert list(preview.values) == list(downloaded.values)
    assert persisted[0]["annual_leave_records"][0]["date"] == "2026-09-03"
    assert extras["employee_values"] == {"E1": 125}


def test_settlement_rolls_back_salaries_when_extras_fail(tmp_path):
    config = DatabaseConfig("sqlite", tmp_path / "rollback.db")
    initialize_database_with_health(config)
    salaries = [_salary("E1"), _salary("E2")]
    for salary in salaries:
        save_employee(config, {"employee_id": salary["employee_id"], "name": salary["employee_id"]})
        save_salary(config, {**salary, "year": 2026, "month": 9}, salary["adjustments"])
    with connect_from_config(config) as conn:
        conn.execute("CREATE TRIGGER reject_extras BEFORE INSERT ON salary_monthly_extras "
                     "BEGIN SELECT RAISE(ABORT, 'extras failed'); END")

    with pytest.raises(sqlite3.IntegrityError, match="extras failed"):
        settle_salary_month(config, 2026, 9, salaries, _extras())

    assert get_settled_month_salaries(config, 2026, 9) == []


@pytest.mark.parametrize("field,value", [
    ("base_salary_snapshot", 33000), ("attendance_bonus_snapshot", 0),
    ("cooling_allowance_snapshot", 0), ("allowance_snapshot", 0),
    ("leave_deduction", 500), ("final_salary", 1),
    ("system_note", "Changed note"), ("manual_note", "Changed manual note"),
])
def test_report_signature_detects_unsaved_payroll_changes(field, value):
    salary = _salary()
    changed = {**salary, field: value}
    assert salary_report_signature(9, [salary], _extras()) != salary_report_signature(9, [changed], _extras())


def test_report_signature_detects_adjustments_and_monthly_extras():
    salary = _salary()
    changed = deepcopy(salary)
    changed["adjustments"][0]["amount"] += 100
    original = salary_report_signature(9, [salary], _extras())
    assert original != salary_report_signature(9, [changed], _extras())
    assert original != salary_report_signature(9, [salary], {**_extras(), "monthly_total": 15000})


def test_monthly_ui_settles_extras_and_blocks_outdated_download(tmp_path):
    from streamlit.testing.v1 import AppTest

    config = DatabaseConfig("sqlite", tmp_path / "ui.db")
    initialize_database_with_health(config)
    salary = _salary()
    save_employee(config, {"employee_id": "E1", "name": "E1"})
    save_salary(config, {**salary, "year": 2026, "month": 9}, salary["adjustments"],
                annual_leave_records=salary["annual_leave_records"])
    app = AppTest.from_string(
        "from pathlib import Path\n"
        "from utils.database import DatabaseConfig\n"
        "from utils.salary_ui import _monthly_tab\n"
        f"_monthly_tab(DatabaseConfig('sqlite', Path({str(config.path)!r})))\n"
    )
    app.session_state["salary_year"] = 2026
    app.session_state["salary_month"] = 9
    app.run()
    assert not app.exception
    app.number_input(key="monthly_extra_previous_2026-09").set_value(14028.0)
    app.number_input(key="monthly_extra_employee_2026-09_E1").set_value(125.0)
    next(button for button in app.button if button.label == "結算薪資").click().run()
    assert not app.exception
    assert get_salary_monthly_extras(config, 2026, 9)["monthly_total"] == 14153
    download = next(item for item in app.get("download_button") if item.proto.label == "下載本月薪資表")
    assert not download.proto.disabled
    app.number_input(key="base_salary_snapshot_2026-09_0").set_value(35000).run()
    assert not app.exception
    download = next(item for item in app.get("download_button") if item.proto.label == "下載本月薪資表")
    assert download.proto.disabled
    assert any("重新按" in warning.value for warning in app.warning)
